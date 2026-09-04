from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from pathlib import Path


OUT = Path(r"C:\Users\Seth\Documents\ChatGPT\AMN Quality\AMN_Trace_Propuesta_Implementacion_MVP.docx")

# Design system: narrative_proposal + proposal_centerpiece
NAVY = "123047"
TEAL = "007B83"
GOLD = "C69214"
INK = "1E2933"
MUTED = "5E6B75"
LINE = "D7E0E5"
PALE_TEAL = "EAF5F5"
PALE_GOLD = "FFF7E5"
WHITE = "FFFFFF"
USABLE = 9360


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right"):
        if edge in kwargs:
            tag = "w:{}".format(edge)
            element = tc_borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tc_borders.append(element)
            for key, value in kwargs[edge].items():
                element.set(qn("w:{}".format(key)), str(value))


def set_cell_margin(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=120):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_layout = tbl_pr.first_child_found_in("w:tblLayout")
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margin(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_font(run, size=None, color=INK, bold=None, italic=None, name="Arial"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_para(p, before=0, after=8, line=1.25, align=None, keep_with_next=False):
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep_with_next
    if align is not None:
        p.alignment = align


def add_run(p, text, **kwargs):
    r = p.add_run(text)
    set_font(r, **kwargs)
    return r


def add_body(doc, text, after=8, bold_lead=None):
    p = doc.add_paragraph(style="Body")
    set_para(p, after=after, line=1.28)
    if bold_lead and text.startswith(bold_lead):
        add_run(p, bold_lead, size=10.5, bold=True)
        add_run(p, text[len(bold_lead):], size=10.5)
    else:
        add_run(p, text, size=10.5)
    return p


def add_heading(doc, text, level=1):
    style = f"Heading {level}"
    p = doc.add_paragraph(style=style)
    if level == 1:
        set_para(p, before=18, after=8, line=1.1, keep_with_next=True)
        add_run(p, text, size=16, color=NAVY, bold=True)
    elif level == 2:
        set_para(p, before=12, after=5, line=1.1, keep_with_next=True)
        add_run(p, text, size=12, color=TEAL, bold=True)
    else:
        set_para(p, before=8, after=3, line=1.1, keep_with_next=True)
        add_run(p, text, size=10.5, color=NAVY, bold=True)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    set_para(p, after=3, line=1.2)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    add_run(p, text, size=10.2)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    set_para(p, after=4, line=1.2)
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.2)
    add_run(p, text, size=10.2)
    return p


def add_callout(doc, label, text, fill=PALE_TEAL):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_repeat_table_header(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, top={"val": "single", "sz": "8", "color": TEAL}, left={"val": "single", "sz": "8", "color": TEAL}, bottom={"val": "single", "sz": "8", "color": TEAL}, right={"val": "single", "sz": "8", "color": TEAL})
    p = cell.paragraphs[0]
    set_para(p, after=0, line=1.15)
    add_run(p, label + " ", size=10.4, color=NAVY, bold=True)
    add_run(p, text, size=10.4)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_table(doc, headers, rows, widths, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, text in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, NAVY)
        set_cell_border(cell, top={"val": "single", "sz": "4", "color": NAVY}, left={"val": "single", "sz": "4", "color": NAVY}, bottom={"val": "single", "sz": "4", "color": NAVY}, right={"val": "single", "sz": "4", "color": NAVY})
        p = cell.paragraphs[0]
        set_para(p, after=0, line=1.0)
        add_run(p, text, size=font_size, color=WHITE, bold=True)
    for r_i, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            cell = cells[idx]
            set_cell_shading(cell, WHITE if r_i % 2 == 0 else "F5F8FA")
            set_cell_border(cell, top={"val": "single", "sz": "2", "color": LINE}, left={"val": "single", "sz": "2", "color": LINE}, bottom={"val": "single", "sz": "2", "color": LINE}, right={"val": "single", "sz": "2", "color": LINE})
            p = cell.paragraphs[0]
            set_para(p, after=0, line=1.1)
            add_run(p, str(text), size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para(p, before=0, after=0, line=1.0)
    add_run(p, "AMN Trace | Propuesta de implementación y MVP | Documento de trabajo", size=8, color=MUTED)


def add_header(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_para(p, before=0, after=0, line=1.0)
    add_run(p, "AMN QUALITY", size=8, color=TEAL, bold=True)
    add_run(p, "  /  AMN TRACE", size=8, color=MUTED)


def build_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.8)
    sec.bottom_margin = Inches(0.7)
    sec.left_margin = Inches(1.0)
    sec.right_margin = Inches(1.0)
    sec.header_distance = Inches(0.3)
    sec.footer_distance = Inches(0.3)
    add_header(sec)
    add_footer(sec)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    for name in ["Body", "Heading 1", "Heading 2", "Heading 3"]:
        if name not in styles:
            styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)

    # Cover / proposal centerpiece
    p = doc.add_paragraph()
    set_para(p, before=72, after=5, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_run(p, "AMN QUALITY", size=11, color=TEAL, bold=True)
    p = doc.add_paragraph()
    set_para(p, before=0, after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_run(p, "AMN Trace", size=30, color=NAVY, bold=True)
    p = doc.add_paragraph()
    set_para(p, before=0, after=24, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.15)
    add_run(p, "Propuesta de implementación y MVP", size=15, color=INK)
    add_run(p, "\nPlataforma de evidencia y trazabilidad de inspección", size=12, color=MUTED)
    line = doc.add_paragraph()
    set_para(line, after=14)
    ppr = line._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "16")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), TEAL)
    border.append(bottom)
    ppr.append(border)
    cover = doc.add_table(rows=2, cols=2)
    set_table_geometry(cover, [4680, 4680])
    set_repeat_table_header(cover.rows[0])
    vals = [("Preparado para", "AMN Quality"), ("Objetivo", "MVP demostrable para trazabilidad de todas las piezas"), ("Horizonte de retención", "Mínimo 5 años"), ("Fecha", "3 de septiembre de 2026")]
    for i, (label, value) in enumerate(vals):
        cell = cover.cell(i // 2, i % 2)
        set_cell_shading(cell, "F5F8FA")
        set_cell_border(cell, top={"val": "single", "sz": "2", "color": LINE}, left={"val": "single", "sz": "2", "color": LINE}, bottom={"val": "single", "sz": "2", "color": LINE}, right={"val": "single", "sz": "2", "color": LINE})
        p = cell.paragraphs[0]
        set_para(p, after=2, line=1.1)
        add_run(p, label + "\n", size=8.5, color=MUTED, bold=True)
        add_run(p, value, size=10.2, color=INK)
    p = doc.add_paragraph()
    set_para(p, before=40, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_run(p, "Documento de trabajo - propuesta inicial", size=9.5, color=MUTED, italic=True)
    doc.add_page_break()

    # Executive overview
    add_heading(doc, "1. Resumen ejecutivo")
    add_body(doc, "AMN Trace es una propuesta para convertir cada inspección visual de AMN Quality en un expediente digital confiable. El sistema asociará la pieza, sus imágenes originales, el resultado de la inspección y el contexto de producción en un registro consultable durante al menos cinco años.")
    add_callout(doc, "Resultado esperado.", "Ante una reclamación o investigación, AMN Quality podrá localizar una pieza por número de serie, lote, orden o fecha y reconstruir en segundos qué ocurrió: dónde se produjo, quién participó, qué máquina la inspeccionó, qué resultado obtuvo y cuál fue la evidencia visual.")
    add_heading(doc, "Problema que resuelve", level=2)
    add_bullet(doc, "La evidencia de inspección suele quedar dispersa en carpetas, equipos o archivos que no garantizan la relación con una pieza específica.")
    add_bullet(doc, "Una falla de red o internet no debe detener la captura de evidencia ni permitir que se pierdan inspecciones.")
    add_bullet(doc, "La evidencia debe permanecer protegida: no se sustituye ni se borra; cualquier corrección se registra como un evento posterior con usuario, fecha y motivo.")
    add_bullet(doc, "Todas las piezas, incluyendo las aceptadas (OK), deben mantener trazabilidad completa a largo plazo.")
    add_heading(doc, "Alcance de esta propuesta", level=2)
    add_body(doc, "El documento define una arquitectura objetivo y un MVP presentable. El MVP valida el flujo crítico: identificar pieza, capturar/recibir evidencia, operar sin internet, sincronizar de forma segura y consultar el expediente desde un portal web.")

    # Requirements
    add_heading(doc, "2. Requisitos de negocio y diseño")
    req_rows = [
        ("Cobertura", "Registrar el 100% de las piezas inspeccionadas: OK, NOK, retrabajo y revisión."),
        ("Retención", "Conservar expediente y evidencia por al menos cinco años, sujeto a la política final de AMN Quality y sus clientes."),
        ("Continuidad", "La estación debe continuar capturando y protegiendo evidencia aunque se interrumpa la red o internet."),
        ("Inmutabilidad", "Imágenes originales y eventos de inspección no se sobrescriben. Se conserva el historial de cambios y acciones."),
        ("Consulta", "Búsqueda por número de parte, serie/DataMatrix, lote, orden, fecha, resultado, línea o estación."),
        ("Arquitectura", "Copia local operativa y repositorio central en nube, con sincronización verificable."),
        ("Seguridad", "Usuarios autenticados, roles, cifrado, bitácora de auditoría y control de acceso por función."),
    ]
    add_table(doc, ["Dimensión", "Requisito"], req_rows, [2250, 7110], font_size=9.4)
    add_heading(doc, "Principios de diseño", level=2)
    add_bullet(doc, "La evidencia original se preserva; los análisis, anotaciones y decisiones se agregan como nuevas capas o eventos.")
    add_bullet(doc, "El identificador de pieza es la llave del expediente; el nombre de archivo sólo facilita su lectura humana.")
    add_bullet(doc, "La nube centraliza la consulta y el respaldo; el edge local mantiene la producción en marcha.")
    add_bullet(doc, "La trazabilidad debe ser verificable, no únicamente visible: integridad, historial y reconciliación son parte del producto.")

    # Architecture
    add_heading(doc, "3. Arquitectura propuesta")
    add_body(doc, "La solución se divide en cuatro capas para separar la operación de planta, la evidencia, la consulta y el resguardo a largo plazo.")
    arch_rows = [
        ("1. Identificación", "Lector de código de barras, QR o DataMatrix; o integración con el identificador que ya emita el proceso."),
        ("2. Edge en planta", "PC industrial o servidor local que recibe imágenes y resultados de la máquina, valida datos, crea el expediente, conserva una cola offline y sincroniza después."),
        ("3. Plataforma central", "API y base de datos para metadatos, usuarios, auditoría, búsquedas y reglas de negocio."),
        ("4. Evidencia en nube", "Almacenamiento cifrado e inmutable para originales, imágenes anotadas y archivos asociados; acceso desde portal web."),
    ]
    add_table(doc, ["Capa", "Responsabilidad"], arch_rows, [2200, 7160])
    add_heading(doc, "Flujo de información", level=2)
    flow = doc.add_table(rows=1, cols=5)
    set_table_geometry(flow, [1800, 1800, 1800, 1800, 2160])
    set_repeat_table_header(flow.rows[0])
    steps = [
        ("1", "Identificar", "Se escanea la pieza."),
        ("2", "Inspeccionar", "Máquina genera imágenes y resultado."),
        ("3", "Cerrar evento", "Edge crea expediente y hash."),
        ("4", "Resguardar", "Copia local y cola offline."),
        ("5", "Sincronizar", "Nube valida, inmoviliza y publica."),
    ]
    for idx, (n, name, desc) in enumerate(steps):
        c = flow.cell(0, idx)
        set_cell_shading(c, PALE_TEAL if idx % 2 == 0 else "F5F8FA")
        set_cell_border(c, top={"val": "single", "sz": "4", "color": TEAL}, left={"val": "single", "sz": "4", "color": TEAL}, bottom={"val": "single", "sz": "4", "color": TEAL}, right={"val": "single", "sz": "4", "color": TEAL})
        p = c.paragraphs[0]
        set_para(p, after=2, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.0)
        add_run(p, n + "\n", size=12, color=TEAL, bold=True)
        add_run(p, name + "\n", size=9.2, color=NAVY, bold=True)
        add_run(p, desc, size=8.2, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    add_callout(doc, "Comportamiento offline.", "Si no existe conectividad, la estación no pierde ni detiene la captura: almacena eventos en una cola local cifrada. Al recuperar conexión, reintenta la sincronización, confirma la integridad de cada archivo y marca el expediente como sincronizado.", fill=PALE_GOLD)

    # Lifecycle and data
    add_heading(doc, "4. Expediente digital de una pieza")
    add_body(doc, "Cada ciclo de inspección genera un evento de trazabilidad. Para que el historial sea comprensible y auditable, se recomienda distinguir entre identidad de pieza, contexto de producción, resultado técnico, archivos de evidencia y eventos de auditoría.")
    fields_rows = [
        ("Identidad", "Número de parte, serie/DataMatrix, ID de evento, lote, orden de producción."),
        ("Contexto", "Planta, línea, estación, máquina/cámara, operador, turno, fecha y hora."),
        ("Resultado", "OK, NOK, retrabajo o revisión; defecto, medidas, reglas/umbrales y receta de inspección."),
        ("Evidencia", "Imagen original, imagen anotada, archivo de medición y hash SHA-256 de cada archivo."),
        ("Configuración", "Versión de algoritmo/modelo, receta, versión de software y parámetros aplicables."),
        ("Auditoría", "Eventos de creación, consulta, exportación, comentario, decisión, corrección y sincronización."),
    ]
    add_table(doc, ["Grupo", "Campos mínimos del MVP"], fields_rows, [2200, 7160])
    add_heading(doc, "Reglas de inmutabilidad y auditoría", level=2)
    add_number(doc, "Al concluir la inspección, el edge calcula un hash para cada archivo original y registra los hashes dentro del evento.")
    add_number(doc, "El registro se cierra y se conserva como evidencia. No se permite reemplazar el archivo original ni editar su resultado sin dejar rastro.")
    add_number(doc, "Si una persona necesita corregir una clasificación o agregar una disposición, crea un evento nuevo enlazado al original, con usuario, sello de tiempo, motivo y comentario.")
    add_number(doc, "Al sincronizar, la plataforma recalcula o verifica los hashes y conserva el objeto en almacenamiento con retención inmutable (WORM) configurada.")
    add_number(doc, "La bitácora de auditoría es de sólo adición: registra quién hizo qué, desde dónde, cuándo y sobre cuál expediente.")

    # MVP Scope
    add_heading(doc, "5. MVP propuesto")
    add_callout(doc, "Meta del MVP.", "Demostrar un circuito completo y creíble de trazabilidad en una estación: una pieza identificada entra al proceso, se inspecciona, deja evidencia local y se consulta desde un portal después de sincronizar.")
    mvp_rows = [
        ("Identificación", "Captura manual o por lector de número de parte y número de serie/DataMatrix.", "Incluido"),
        ("Captura de inspección", "Recepción de resultado e imagen desde carpeta, API, archivo exportado o conector simulado de la máquina.", "Incluido"),
        ("Expediente", "Creación de registro, metadatos, estado OK/NOK y asociación con evidencia.", "Incluido"),
        ("Edge offline", "Persistencia local, cola de sincronización y reintento cuando se restablece la red.", "Incluido"),
        ("Portal", "Búsqueda, ficha de pieza, visualización de imágenes y línea de tiempo de auditoría.", "Incluido"),
        ("Integridad", "Hash de archivo, control de cambios y retención lógica de evidencia.", "Incluido"),
        ("ERP/MES", "Lectura automática de órdenes, bloqueo de producción o intercambio bidireccional.", "Fase posterior"),
        ("Analítica avanzada", "Indicadores, alertas de tendencias y modelos de IA adicionales.", "Fase posterior"),
        ("Multi-planta", "Gestión de varias plantas, redes segregadas o gran volumen empresarial.", "Fase posterior"),
    ]
    add_table(doc, ["Módulo", "Capacidad", "MVP"], mvp_rows, [1900, 5700, 1760], font_size=8.9)
    add_heading(doc, "Demostración sugerida para presentación", level=2)
    demo_rows = [
        ("01", "Escanear", "Introducir o leer el ID de una pieza de ejemplo."),
        ("02", "Inspeccionar", "Generar una inspección OK y otra NOK con imágenes de muestra."),
        ("03", "Interrumpir red", "Mostrar que los eventos quedan en la cola local sin pérdida."),
        ("04", "Recuperar red", "Sincronizar y confirmar que hashes y archivos coinciden."),
        ("05", "Consultar", "Buscar una serie y abrir su expediente con imagen, resultado, datos de producción y auditoría."),
    ]
    add_table(doc, ["Paso", "Acción", "Lo que demuestra"], demo_rows, [800, 2000, 6560])

    # Technical approach
    add_heading(doc, "6. Diseño técnico recomendado para el MVP")
    tech_rows = [
        ("Aplicación edge", "Servicio local Windows/Linux con base de datos embebida o local, carpeta de evidencia cifrada y agente de sincronización."),
        ("Integración con máquina", "Conector adaptador según capacidad real: API, carpeta vigilada, archivo CSV/XML/JSON, base de datos o protocolo del fabricante."),
        ("API central", "Servicio seguro para expedientes, autenticación, colas de carga, metadatos, búsqueda y bitácora."),
        ("Base de datos", "Base relacional para relaciones y consultas: pieza, lote, orden, estación, inspección, evidencia, usuario y auditoría."),
        ("Evidencia", "Almacenamiento de objetos para imágenes/archivos, con cifrado, versionado y política WORM/retención inmutable."),
        ("Portal web", "Interfaz de consulta con filtros, ficha de pieza, visor de imagen, hashes, historial y exportación controlada."),
        ("Observabilidad", "Monitoreo de cola, capacidad local, fallas de sincronización, integridad y salud de estación."),
    ]
    add_table(doc, ["Componente", "Propuesta"], tech_rows, [2300, 7060])
    add_heading(doc, "Convención de archivos", level=2)
    add_body(doc, "El sistema debe guardar una clave técnica única y puede utilizar un nombre legible para operación diaria. Ejemplo: ", after=4)
    p = doc.add_paragraph()
    set_para(p, before=0, after=10, line=1.1)
    p.paragraph_format.left_indent = Inches(0.3)
    add_run(p, "AMN_[parte]_[serie]_[resultado]_[fecha-hora]_[cámara].jpg", size=10, color=NAVY, bold=True, name="Courier New")
    add_body(doc, "La asociación válida se conserva en la base de datos mediante ID de evidencia, hash y referencias al expediente; el nombre del archivo no sustituye esa relación.")

    # Pilot delivery plan
    add_heading(doc, "7. Plan de implementación")
    phases = [
        ("0. Descubrimiento", "Validar datos de la máquina, identificación de pieza, volumen, imágenes por ciclo y reglas de acceso.", "2-5 días hábiles"),
        ("1. Diseño del MVP", "Definir integración, modelo de datos, infraestructura inicial, casos de prueba y demostración.", "3-5 días hábiles"),
        ("2. Construcción", "Implementar edge, API/portal, almacenamiento y bitácora para una estación piloto.", "2-4 semanas"),
        ("3. Pruebas piloto", "Probar operación normal, red caída, reinicio de estación, recuperación y búsquedas de auditoría.", "1-2 semanas"),
        ("4. Validación y siguiente fase", "Documentar hallazgos, estimar escalamiento y definir integración con ERP/MES si aplica.", "1 semana"),
    ]
    add_table(doc, ["Fase", "Resultado", "Estimación inicial"], phases, [1800, 5660, 1900])
    add_callout(doc, "Nota de planeación.", "Las estimaciones se afinan después de conocer la interfaz de la máquina de inspección y el volumen real. La mayor variable técnica suele ser la forma en que cada máquina entrega imágenes y resultados.", fill=PALE_GOLD)
    add_heading(doc, "Criterios de aceptación del piloto", level=2)
    for item in [
        "Se registra una pieza OK y una NOK con sus imágenes originales y metadatos completos.",
        "Se puede buscar cada pieza por serie/ID, número de parte y rango de fecha.",
        "La estación sigue creando expedientes cuando no hay conectividad.",
        "Al restaurarse la red, los eventos pendientes se sincronizan sin duplicados y con verificación de hash.",
        "Una corrección queda registrada como evento nuevo y no destruye el registro original.",
        "Un usuario autorizado visualiza la línea de tiempo y la evidencia sin poder alterar originales.",
    ]:
        add_bullet(doc, item)

    # Sizing and risks
    add_heading(doc, "8. Dimensionamiento, riesgos y decisiones pendientes")
    add_heading(doc, "Estimación de almacenamiento", level=2)
    add_body(doc, "Antes de contratar capacidad definitiva se debe medir el volumen de un turno representativo. La fórmula inicial es:")
    p = doc.add_paragraph()
    set_para(p, after=9, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.1)
    add_run(p, "Piezas/día x imágenes/pieza x MB/imagen x días/año x años de retención", size=10.2, color=NAVY, bold=True)
    add_body(doc, "A ese resultado se agregan márgenes para copias de seguridad, imágenes anotadas, archivos de medición, crecimiento de producción y retención local. Como decisión práctica, el almacenamiento local puede conservar la ventana operativa reciente (por ejemplo, 30 a 180 días) y la nube mantener los cinco años, salvo que el cliente pida el mismo horizonte en ambos sitios.")
    risks = [
        ("Integración", "No conocer aún cómo la máquina exporta imágenes/resultados.", "Solicitar manual, muestras de archivos y acceso de prueba; diseñar un adaptador desacoplado."),
        ("Identidad", "Piezas sin número de serie único o lectura inconsistente.", "Definir la llave de trazabilidad y el punto del proceso donde se lee/genera."),
        ("Volumen", "Subestimar tamaño de imágenes y crecimiento de cinco años.", "Medir una semana representativa antes de cerrar costos de nube y hardware."),
        ("Operación", "Intervenciones manuales que rompan el vínculo pieza-evidencia.", "Validaciones en pantalla, roles y eventos de excepción con motivo obligatorio."),
        ("Cumplimiento", "Reglas específicas del cliente final no confirmadas.", "Alinear retención, exportación, aprobaciones y evidencias contra requisitos vigentes del cliente."),
    ]
    add_table(doc, ["Riesgo", "Qué falta confirmar", "Mitigación"], risks, [1700, 3800, 3860], font_size=8.7)
    add_heading(doc, "Información necesaria para iniciar", level=2)
    for item in [
        "Marca/modelo de la máquina y documentación de integración disponible.",
        "Ejemplos reales anonimizados de imágenes, resultados y archivos exportados.",
        "Volumen por día, imágenes por pieza, tamaño promedio y número de estaciones.",
        "Cómo se genera o lee el número de serie/DataMatrix y en qué momento del proceso.",
        "Lista inicial de usuarios/roles: operador, calidad, supervisor, auditor y administrador.",
        "Políticas de red, preferencias de nube, requisitos del cliente final y regla de retención final.",
    ]:
        add_bullet(doc, item)

    # Closing
    add_heading(doc, "9. Próximo paso recomendado")
    add_callout(doc, "Propuesta inmediata.", "Realizar una sesión técnica de 60-90 minutos con producción, calidad e IT para confirmar el método de integración de una estación. Con esa información se prepara el diseño de piloto, la demostración del MVP y una estimación de infraestructura basada en datos reales.")
    add_body(doc, "El MVP debe centrarse en probar la confianza del proceso, no en cubrir todas las integraciones corporativas desde el primer día. Una estación, un flujo de identificación, un conector de inspección, operación offline y consulta central son suficientes para demostrar el valor de AMN Trace y reducir el riesgo de la siguiente fase.")
    p = doc.add_paragraph()
    set_para(p, before=24, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_run(p, "AMN Trace: cada pieza deja evidencia; cada evidencia conserva su historia.", size=11, color=TEAL, bold=True, italic=True)

    doc.core_properties.title = "AMN Trace - Propuesta de implementación y MVP"
    doc.core_properties.subject = "Trazabilidad industrial y evidencia de inspección"
    doc.core_properties.author = "AMN Quality"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_doc()
